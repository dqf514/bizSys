"""批次3：列表导出补齐（OA 我的申请/审批中心/部门/审计日志/字典）+ 记录时间修正（留痕/权限/校验）。"""

import io
import json
import uuid
from datetime import datetime

from openpyxl import load_workbook
from sqlalchemy import select

from app.modules.admin.models import AuditLog
from app.modules.approval.models import ApprovalInstance, ApprovalRecord
from app.modules.finance.models import Settlement
from app.modules.report.models import Report

UUID = uuid.uuid4().hex[:8]


def _login(client, username: str, password: str) -> dict:
    r = client.post("/api/v1/auth/login", json={"username": username, "password": password})
    data = r.json()
    assert data["code"] == 0, data
    return {"Authorization": f"Bearer {data['data']['access_token']}"}


def _no_perm_user(client, admin_headers) -> dict:
    suf = uuid.uuid4().hex[:8]
    r = client.post(
        "/api/v1/admin/users",
        headers=admin_headers,
        json={"username": f"noperm_{suf}", "password": "Test@123456", "name": "无权限", "status": "enabled",
              "role_ids": []},
    ).json()
    assert r["code"] == 0, r
    return _login(client, f"noperm_{suf}", "Test@123456")


def _xlsx_rows(resp) -> list[list]:
    assert resp.status_code == 200, resp.content[:200]
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    return [[c.value for c in row] for row in ws.iter_rows()]


# ---------------- 导出补齐 ----------------


def test_oa_my_export(client, admin_headers):
    """发起一条请假申请后，我的申请导出含该单（状态/分类已转中文标签）。"""
    catalog = client.get("/api/v1/oa/catalog", headers=admin_headers).json()["data"]
    form = next(f for g in catalog for f in g["forms"] if f["key"] == "leave_application")
    suf = uuid.uuid4().hex[:8]
    r = client.post(
        "/api/v1/oa/instances",
        headers=admin_headers,
        json={"formId": form["id"],
              "data": {"leave_type": "事假", "date_range": ["2026-02-01", "2026-02-02"], "reason": f"导出测试_{suf}"}},
    ).json()
    assert r["code"] == 0, r
    display_no = r["data"]["displayNo"]

    rows = _xlsx_rows(client.get("/api/v1/oa/instances/my/export", headers=admin_headers))
    assert rows[0][:4] == ["申请编号", "表单", "分类", "状态"]
    hit = [row for row in rows[1:] if row[0] == display_no]
    assert hit and hit[0][2] == "请假" and hit[0][3] == "处理中"


def test_approval_tasks_export(client, admin_headers):
    """审批中心三个页签均可导出 xlsx。"""
    for tab in ("todo", "done", "initiated"):
        rows = _xlsx_rows(client.get("/api/v1/approval/tasks/export", headers=admin_headers, params={"tab": tab}))
        assert rows[0][:2] == ["标题", "业务类型"]
        assert len(rows) >= 1
    # initiated 页签应有种子/前置用例发起的实例
    rows = _xlsx_rows(client.get("/api/v1/approval/tasks/export", headers=admin_headers, params={"tab": "initiated"}))
    assert len(rows) >= 2


def test_admin_exports(client, admin_headers):
    """部门 / 字典 / 审计日志导出（沿用列表查看权限码）。"""
    rows = _xlsx_rows(client.get("/api/v1/admin/depts/export", headers=admin_headers))
    assert rows[0] == ["部门名称", "部门编码", "上级部门", "排序", "状态"]
    assert any(row[1] == "D00" and not row[2] for row in rows[1:])  # openpyxl 空单元格读出 None
    assert any(row[1] == "D01" and row[2] == "总经办" for row in rows[1:])

    rows = _xlsx_rows(client.get("/api/v1/admin/dict/types/export", headers=admin_headers))
    assert rows[0] == ["类型编码", "类型名称", "标签", "英文标签", "字典值", "排序", "状态"]
    assert any(row[0] == "oa_category" and row[4] == "leave" for row in rows[1:])

    rows = _xlsx_rows(client.get("/api/v1/admin/audit-logs/export", headers=admin_headers))
    assert rows[0][:3] == ["操作人", "模块", "动作"]
    assert len(rows) >= 2  # 前面用例已产生审计记录


def test_export_perm_guard(client, admin_headers):
    """无权限用户：部门/字典/审计导出 403；OA/审批导出与列表同权。"""
    headers = _no_perm_user(client, admin_headers)
    for url in ("/api/v1/admin/depts/export", "/api/v1/admin/dict/types/export", "/api/v1/admin/audit-logs/export"):
        resp = client.get(url, headers=headers)
        assert resp.status_code == 403 and resp.json()["code"] == 20002, url
    resp = client.get("/api/v1/oa/instances/my/export", headers=headers)
    assert resp.status_code == 403 and resp.json()["code"] == 20002
    # 审批中心导出与列表同为登录可用
    resp = client.get("/api/v1/approval/tasks/export", headers=headers, params={"tab": "todo"})
    assert resp.status_code == 200


# ---------------- 记录时间修正 ----------------


def _make_report(db) -> Report:
    r = Report(
        report_no=f"BG-RT-{uuid.uuid4().hex[:8]}",
        task_id=0,
        task_item_id=0,
        file_id=0,
        status="sealed",
        seal_time=datetime(2026, 1, 10, 9, 0, 0),
    )
    db.add(r)
    db.commit()
    return r


def test_retime_report_flow(client, admin_headers, db):
    """报告创建/签章时间修正：lookup → 改值 → 字段更新 + sys_audit_log 留痕（原值/新值/原因）。"""
    report = _make_report(db)

    r = client.get(
        "/api/v1/admin/maintenance/retime/lookup",
        headers=admin_headers,
        params={"bizType": "report", "key": report.report_no},
    ).json()
    assert r["code"] == 0, r
    assert r["data"]["id"] == report.id
    fields = {f["field"]: f for f in r["data"]["fields"]}
    assert set(fields) == {"create_time", "seal_time"}
    assert fields["seal_time"]["value"] == "2026-01-10 09:00:00"

    r = client.post(
        "/api/v1/admin/maintenance/retime",
        headers=admin_headers,
        json={"bizType": "report", "id": report.id, "field": "seal_time",
              "value": "2026-01-12 15:30:00", "reason": "签章系统时钟异常，按纸质记录修正"},
    ).json()
    assert r["code"] == 0, r
    assert r["data"]["value"] == "2026-01-12 15:30:00"

    db.expire_all()
    assert db.get(Report, report.id).seal_time == datetime(2026, 1, 12, 15, 30, 0)

    log = db.scalar(
        select(AuditLog)
        .where(AuditLog.action == "maintenance_retime", AuditLog.biz_type == "report", AuditLog.biz_id == report.id)
        .order_by(AuditLog.id.desc())
    )
    assert log is not None
    detail = json.loads(log.detail)
    assert detail["field"] == "seal_time"
    assert detail["old"] == "2026-01-10 09:00:00"
    assert detail["new"] == "2026-01-12 15:30:00"
    assert detail["reason"] == "签章系统时钟异常，按纸质记录修正"


def test_retime_approval_record_and_settlement(client, admin_headers, db):
    """审批记录 at / 结算单 create_time 修正（按 ID / 结算单号检索）。"""
    inst = ApprovalInstance(biz_type="test", biz_id=0, title="修正测试实例", initiator_id=1, initiator_name="测试")
    db.add(inst)
    db.commit()
    rec = ApprovalRecord(instance_id=inst.id, seq=0, node_name="测试节点", action="agree", operator_id=1,
                         operator_name="测试", at=datetime(2026, 1, 5, 8, 0, 0))
    db.add(rec)
    db.commit()
    r = client.get(
        "/api/v1/admin/maintenance/retime/lookup",
        headers=admin_headers,
        params={"bizType": "approval_record", "key": str(rec.id)},
    ).json()
    assert r["code"] == 0 and [f["field"] for f in r["data"]["fields"]] == ["at"]
    r = client.post(
        "/api/v1/admin/maintenance/retime",
        headers=admin_headers,
        json={"bizType": "approval_record", "id": rec.id, "field": "at",
              "value": "2026-01-05 10:00:00", "reason": "补录审批时间"},
    ).json()
    assert r["code"] == 0, r
    db.expire_all()
    assert db.get(ApprovalRecord, rec.id).at == datetime(2026, 1, 5, 10, 0, 0)

    st = Settlement(settlement_no=f"HK-RT-{uuid.uuid4().hex[:8]}", task_id=0, task_no="", amount=1,
                    invoiceable_amount=1, status="pending")
    db.add(st)
    db.commit()
    r = client.get(
        "/api/v1/admin/maintenance/retime/lookup",
        headers=admin_headers,
        params={"bizType": "settlement", "key": st.settlement_no},
    ).json()
    assert r["code"] == 0 and [f["field"] for f in r["data"]["fields"]] == ["create_time"]
    r = client.post(
        "/api/v1/admin/maintenance/retime",
        headers=admin_headers,
        json={"bizType": "settlement", "id": st.id, "field": "create_time",
              "value": "2026-01-03 09:00:00", "reason": "补登记"},
    ).json()
    assert r["code"] == 0, r
    db.expire_all()
    assert db.get(Settlement, st.id).create_time == datetime(2026, 1, 3, 9, 0, 0)


def test_retime_validation_and_perm(client, admin_headers, manager_headers, db):
    """原因必填 / 非法字段 / 未知类型 / 单据不存在 / 非超管拦截。"""
    report = _make_report(db)
    base = {"bizType": "report", "id": report.id, "field": "create_time", "value": "2026-01-02 00:00:00"}

    r = client.post("/api/v1/admin/maintenance/retime", headers=admin_headers, json=base).json()
    assert r["code"] == 10001  # 缺 reason
    r = client.post("/api/v1/admin/maintenance/retime", headers=admin_headers,
                    json={**base, "reason": "x", "value": "2026/01/02"}).json()
    assert r["code"] == 10001  # 时间格式非法
    r = client.post("/api/v1/admin/maintenance/retime", headers=admin_headers,
                    json={**base, "reason": "x", "field": "status"}).json()
    assert r["code"] == 10001  # 非白名单字段
    r = client.post("/api/v1/admin/maintenance/retime", headers=admin_headers,
                    json={**base, "reason": "x", "bizType": "user"}).json()
    assert r["code"] == 10001  # 未知业务类型
    r = client.get("/api/v1/admin/maintenance/retime/lookup", headers=admin_headers,
                   params={"bizType": "report", "key": "BG-NOT-EXIST"}).json()
    assert r["code"] == 10002

    # biz_user 无 admin:maintenance:retime → 403
    r = client.get("/api/v1/admin/maintenance/retime/lookup", headers=manager_headers,
                   params={"bizType": "report", "key": report.report_no}).json()
    assert r["code"] == 20002
    r = client.post("/api/v1/admin/maintenance/retime", headers=manager_headers,
                    json={**base, "reason": "x"}).json()
    assert r["code"] == 20002
