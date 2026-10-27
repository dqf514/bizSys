"""批次A：导航整合（业务系统分组/样品权限）+ 跨任务样品列表与导出。"""

import io
import uuid

from openpyxl import load_workbook

UUID = uuid.uuid4().hex[:8]


def _login(client, username: str, password: str) -> dict:
    r = client.post("/api/v1/auth/login", json={"username": username, "password": password})
    data = r.json()
    assert data["code"] == 0, data
    return {"Authorization": f"Bearer {data['data']['access_token']}"}


def _create_user(client, admin_headers, role_ids: list[int]) -> dict:
    suf = uuid.uuid4().hex[:8]
    r = client.post(
        "/api/v1/admin/users",
        headers=admin_headers,
        json={"username": f"ba_{suf}", "password": "Test@123456", "name": "批次A", "status": "enabled",
              "role_ids": role_ids},
    ).json()
    assert r["code"] == 0, r
    return _login(client, f"ba_{suf}", "Test@123456")


def _make_task_with_sample(client, admin_headers) -> int:
    r = client.post(
        "/api/v1/task/tasks",
        headers=admin_headers,
        json={
            "taskName": f"ba_样品任务_{UUID}",
            "samples": [{"name": f"ba_样品_{UUID}", "model": "M1", "vin": f"VIN{UUID}", "vehicleType": "轿车"}],
        },
    ).json()
    assert r["code"] == 0, r
    return r["data"]["id"]


# ---------------- 样品列表 / 导出 ----------------


def test_sample_list_and_search(client, admin_headers):
    """跨任务样品列表：返回任务/客户字段，q 可按 VIN 检索。"""
    tid = _make_task_with_sample(client, admin_headers)
    r = client.get("/api/v1/task/samples", headers=admin_headers, params={"q": f"VIN{UUID}"})
    body = r.json()
    assert body["code"] == 0, body
    items = body["data"]["items"]
    assert len(items) == 1, items
    row = items[0]
    assert row["taskId"] == tid and row["taskName"] == f"ba_样品任务_{UUID}"
    assert row["name"] == f"ba_样品_{UUID}" and row["inStock"] == "in"
    # 状态页签
    r = client.get("/api/v1/task/samples", headers=admin_headers, params={"status": "out", "q": f"VIN{UUID}"})
    assert r.json()["data"]["total"] == 0
    r = client.get("/api/v1/task/samples/counts", headers=admin_headers)
    assert r.json()["code"] == 0 and "in" in r.json()["data"]


def test_sample_list_pagination_bound(client, admin_headers):
    """page 上限同样约束样品列表（复用 PAGE_MAX）。"""
    r = client.get("/api/v1/task/samples", headers=admin_headers, params={"page": 10**19})
    assert r.status_code == 422 and r.json()["code"] == 10001


def test_sample_export(client, admin_headers):
    """样品导出 xlsx：表头中文、含检索到的行。"""
    _make_task_with_sample(client, admin_headers)
    r = client.post("/api/v1/task/samples/export", headers=admin_headers)
    assert r.status_code == 200 and r.content[:2] == b"PK"
    wb = load_workbook(io.BytesIO(r.content))
    rows = [[c.value for c in row] for row in wb.active.iter_rows()]
    assert rows[0][:3] == ["样品编号", "样品名称", "型号"]
    assert any(row[1] == f"ba_样品_{UUID}" for row in rows[1:])


def test_sample_perm_guard(client, admin_headers):
    """task:sample:list 权限：无权限 403；业务角色（种子含该权限点）200。"""
    no_perm = _create_user(client, admin_headers, [])
    r = client.get("/api/v1/task/samples", headers=no_perm)
    assert r.status_code == 403
    # 种子 biz_user 角色含 task:sample:list
    r = client.get("/api/v1/task/samples", headers=manager_headers(client))
    assert r.json()["code"] == 0


def manager_headers(client):
    return _login(client, "manager1", "Test@123456")


# ---------------- 菜单结构（种子） ----------------


def test_menu_tree_biz_group(client, admin_headers):
    """业务系统分组含客户/合同/任务/报告/样品/财务；流程中心隐藏但仍挂权限点。

    测试库跨运行复用、种子不回写 parent_id/sort，先执行批次E迁移脚本（幂等）再断言目标结构。
    """
    from scripts.migrate_nav_batch_e import main as migrate_nav

    migrate_nav()
    tree = client.get("/api/v1/admin/menus/tree", headers=admin_headers).json()["data"]

    def find(nodes, name):
        for n in nodes:
            if n["name"] == name:
                return n
            r = find(n.get("children") or [], name)
            if r:
                return r
        return None

    biz = find(tree, "BizSystem")
    assert biz and biz["menuType"] == "dir", biz
    child_names = [c["name"] for c in biz["children"]]
    expected = ["Customers", "Contracts", "Tasks", "Reports", "Samples", "Procurement", "Finance"]
    assert child_names == expected, child_names
    sample = find(tree, "Samples")
    assert sample["path"] == "/task/samples" and sample["permission"] == "task:sample:list"
    finance = find(tree, "Finance")
    fin_names = [c["name"] for c in finance["children"]]
    assert fin_names == ["FinanceInvoices", "FinanceSettlements", "FinancePayments"], fin_names
    approval = find(tree, "ApprovalCenter")
    assert approval and approval["visible"] == 0  # 隐藏但保留（approval:view 权限载体）

    # /auth/me：导航树不含隐藏的流程中心，allMenuPaths 包含 /approval（深链）
    me = client.get("/api/v1/auth/me", headers=admin_headers).json()["data"]
    nav_names = [n["name"] for n in me["menus"]]
    assert "ApprovalCenter" not in nav_names
    assert "/approval" in me["allMenuPaths"]
