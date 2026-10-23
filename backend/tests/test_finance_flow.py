"""财务主链路（HTTP）：任务 finished → 结算单(HK) → 开票审批 → 发票登记 → 到账(DZ) → paid → 导出。"""

import io
import uuid

from openpyxl import load_workbook

UUID = uuid.uuid4().hex[:8]


def _login(client, username: str, password: str) -> dict:
    r = client.post("/api/v1/auth/login", json={"username": username, "password": password})
    data = r.json()
    assert data["code"] == 0, data
    return {"Authorization": f"Bearer {data['data']['access_token']}"}


def _approved_contract(client, headers, amount=10000):
    suf = uuid.uuid4().hex[:8]  # 每次调用独立后缀：测试库数据跨用例持久
    cust = client.post(
        "/api/v1/crm/customers",
        headers=headers,
        json={"name": f"财务客户_{suf}", "tax_no": f"913401FIN{suf}", "status": "enabled"},
    ).json()
    assert cust["code"] == 0, cust
    con = client.post(
        "/api/v1/contract/contracts",
        headers=headers,
        json={"name": f"财务合同_{suf}", "customer_id": cust["data"]["id"], "amount": amount},
    ).json()
    cid = con["data"]["id"]
    assert client.post(f"/api/v1/contract/contracts/{cid}/submit", headers=headers, json={}).json()["code"] == 0

    from sqlalchemy import select

    from app.core.db import SessionLocal
    from app.modules.approval import service as approval
    from app.modules.approval.models import ApprovalInstance, ApprovalTask
    from app.modules.auth.models import User

    db = SessionLocal()
    try:
        for _ in range(2):  # 两级依次同意
            target = None
            for t in db.scalars(
                select(ApprovalTask).where(ApprovalTask.status == "pending", ApprovalTask.deleted == 0)
            ):
                inst = db.get(ApprovalInstance, t.instance_id)
                if inst and inst.biz_type == "contract" and inst.biz_id == cid:
                    target = t
                    break
            assert target, "contract instance task not found"
            approval.approve(db, target.id, db.get(User, target.assignee_id), "ok")
            db.commit()
    finally:
        db.close()
    return cid


def _finished_task(client, headers):
    cid = _approved_contract(client, headers)
    suf = uuid.uuid4().hex[:8]
    r = client.post(
        "/api/v1/task/tasks",
        headers=headers,
        json={"task_name": f"财务任务_{suf}", "contract_id": cid},
    ).json()
    assert r["code"] == 0, r
    tid = r["data"]["id"]
    for action, expect in [
        ("submit", "in_progress"),
        ("finish", "finished"),
    ]:
        r = client.post(f"/api/v1/task/tasks/{tid}/{action}", headers=headers).json()
        assert r["code"] == 0 and r["data"]["status"] == expect, (action, r)
    return tid, r["data"] and client.get(f"/api/v1/task/tasks/{tid}", headers=headers).json()["data"]


def _todo_id(client, headers, biz_type: str, biz_id: int) -> int:
    page = client.get("/api/v1/approval/tasks", headers=headers, params={"tab": "todo", "size": 100}).json()
    todo = [i for i in page["data"]["items"] if i["bizType"] == biz_type and i["bizId"] == biz_id]
    assert todo, f"no todo for {biz_type}/{biz_id}"
    return todo[0]["id"]


def _confirm_settlement(client, headers, sid: int) -> None:
    """发送客户确认 + 代客标记已确认（开票前置）。"""
    opts = client.get("/api/v1/admin/users/options", headers=headers).json()["data"]
    r = client.post(
        f"/api/v1/finance/settlements/{sid}/send-confirm",
        headers=headers,
        json={"receiver_id": opts[0]["id"], "note": "请与客户核对"},
    ).json()
    assert r["code"] == 0 and r["data"]["confirmStatus"] == "pending", r
    r = client.post(
        f"/api/v1/finance/settlements/{sid}/confirm",
        headers=headers,
        json={"action": "confirm", "note": "客户已确认"},
    ).json()
    assert r["code"] == 0 and r["data"]["confirmStatus"] == "confirmed", r


def test_finance_full_flow(client, admin_headers, manager_headers, leader_headers):
    headers = admin_headers
    tid, task = _finished_task(client, headers)

    # task-options 可见（finished 且无结算单），格式 {value,label,sub}
    opts = client.get("/api/v1/finance/settlements/task-options", headers=headers).json()["data"]
    opt = [o for o in opts if o["value"] == tid]
    assert opt and opt[0]["label"] == task["taskNo"] and task["customerName"] in opt[0]["sub"]

    # 非任务创建人（leader1 有 finance:settlement:create 但非创建人/非超管）→ 80003
    r = client.post(
        "/api/v1/finance/settlements", headers=leader_headers, json={"task_id": tid, "amount": 10000}
    ).json()
    assert r["code"] == 80003, r

    # 创建结算单：HK 编号、invoiceableAmount 默认=amount、任务 finished → settled
    r = client.post(
        "/api/v1/finance/settlements",
        headers=headers,
        json={"task_id": tid, "amount": 10000, "remark": "首款结算"},
    ).json()
    assert r["code"] == 0, r
    sid = r["data"]["id"]
    assert r["data"]["settlementNo"].startswith("HK")
    assert r["data"]["status"] == "pending"
    assert r["data"]["invoiceableAmount"] == 10000.0
    t = client.get(f"/api/v1/task/tasks/{tid}", headers=headers).json()["data"]
    assert t["status"] == "settled"

    # 同一任务重复创建 → 80002
    r = client.post("/api/v1/finance/settlements", headers=headers, json={"task_id": tid, "amount": 1}).json()
    assert r["code"] == 80002

    # task-options 不再包含该任务
    opts = client.get("/api/v1/finance/settlements/task-options", headers=headers).json()["data"]
    assert all(o["value"] != tid for o in opts)

    # 未开票（pending）登记到账 → 80001
    r = client.post(
        "/api/v1/finance/payments",
        headers=headers,
        json={"settlement_id": sid, "pay_type": "bank_transfer", "payer_name": "X", "amount": 1,
              "paid_at": "2026-10-09"},
    ).json()
    assert r["code"] == 80001

    # 未发送客户确认：直接标记确认 → 80001；申请开票（未确认）→ 80001
    r = client.post(
        f"/api/v1/finance/settlements/{sid}/confirm", headers=headers, json={"action": "confirm", "note": "x"}
    ).json()
    assert r["code"] == 80001
    r = client.post(f"/api/v1/finance/settlements/{sid}/apply-invoice", headers=headers, json={}).json()
    assert r["code"] == 80001

    # 发送客户确认 + 代客确认后可申请开票
    _confirm_settlement(client, headers, sid)

    # 申请开票 → invoicing；审批驳回 → 回 pending；再申请 → 通过 → invoiced + 发票行
    r = client.post(f"/api/v1/finance/settlements/{sid}/apply-invoice", headers=headers, json={}).json()
    assert r["code"] == 0 and r["data"]["status"] == "invoicing", r
    todo = _todo_id(client, manager_headers, "invoice", sid)
    r = client.post(
        f"/api/v1/approval/tasks/{todo}/reject", headers=manager_headers, json={"comment": "退回补充"}
    ).json()
    assert r["code"] == 0, r
    s = client.get(f"/api/v1/finance/settlements/{sid}", headers=headers).json()["data"]
    assert s["status"] == "pending"

    r = client.post(f"/api/v1/finance/settlements/{sid}/apply-invoice", headers=headers, json={}).json()
    assert r["code"] == 0 and r["data"]["status"] == "invoicing", r
    todo = _todo_id(client, manager_headers, "invoice", sid)
    r = client.post(f"/api/v1/approval/tasks/{todo}/approve", headers=manager_headers, json={"comment": "同意"}).json()
    assert r["code"] == 0, r
    s = client.get(f"/api/v1/finance/settlements/{sid}", headers=headers).json()["data"]
    assert s["status"] == "invoiced"
    assert len(s["invoices"]) == 1 and s["invoices"][0]["status"] == "issued"
    inv = s["invoices"][0]
    assert inv["invoiceNo"] == ""  # 待财务登记

    # 财务登记发票信息
    r = client.put(
        f"/api/v1/finance/invoices/{inv['id']}",
        headers=headers,
        json={"invoice_no": f"INV-{UUID}", "amount": 10000, "invoice_date": "2026-10-09"},
    ).json()
    assert r["code"] == 0 and r["data"]["invoiceNo"] == f"INV-{UUID}", r

    # 到账登记两笔（DZ 编号；累计达到金额 → paid）
    r = client.post(
        "/api/v1/finance/payments",
        headers=headers,
        json={
            "settlement_id": sid,
            "pay_type": "bank_transfer",
            "payer_name": "测试付款方",
            "payer_account": "6222020200112233",
            "amount": 4000,
            "bank": "工商银行",
            "paid_at": "2026-10-09",
        },
    ).json()
    assert r["code"] == 0 and r["data"]["paymentNo"].startswith("DZ"), r
    assert r["data"]["settlementNo"] == s["settlementNo"]
    s = client.get(f"/api/v1/finance/settlements/{sid}", headers=headers).json()["data"]
    assert s["status"] == "invoiced" and s["paidAmount"] == 4000.0

    r = client.post(
        "/api/v1/finance/payments",
        headers=headers,
        json={"settlement_id": sid, "pay_type": "draft", "payer_name": "测试付款方", "amount": 6000,
              "paid_at": "2026-10-10"},
    ).json()
    assert r["code"] == 0, r
    s = client.get(f"/api/v1/finance/settlements/{sid}", headers=headers).json()["data"]
    assert s["status"] == "paid"
    assert len(s["payments"]) == 2

    # 列表检索
    lst = client.get("/api/v1/finance/settlements", headers=headers, params={"q": s["settlementNo"]}).json()["data"]
    assert lst["total"] == 1 and lst["items"][0]["status"] == "paid"
    pays = client.get("/api/v1/finance/payments", headers=headers, params={"q": "测试付款方"}).json()["data"]
    assert pays["total"] >= 2 and pays["items"][0]["settlementNo"]

    # 导出（GET，xlsx 流）
    for url in ("/api/v1/finance/settlements/export", "/api/v1/finance/payments/export"):
        resp = client.get(url, headers=headers)
        assert resp.status_code == 200, url
        wb = load_workbook(io.BytesIO(resp.content))
        assert wb.active.max_row >= 2


def test_finance_perm_forbidden(client, admin_headers):
    """无权限角色访问财务接口 → 403（20002）。"""
    suf = uuid.uuid4().hex[:8]
    username = f"nofin_{suf}"
    r = client.post(
        "/api/v1/admin/users",
        headers=admin_headers,
        json={"username": username, "password": "Test@123456", "name": "无财务权限", "status": "enabled",
              "role_ids": []},
    ).json()
    assert r["code"] == 0, r
    headers = _login(client, username, "Test@123456")

    resp = client.get("/api/v1/finance/settlements", headers=headers)
    assert resp.status_code == 403 and resp.json()["code"] == 20002
    resp = client.get("/api/v1/finance/invoices", headers=headers)
    assert resp.status_code == 403 and resp.json()["code"] == 20002
    resp = client.get("/api/v1/finance/invoices/export", headers=headers)
    assert resp.status_code == 403 and resp.json()["code"] == 20002
    resp = client.get("/api/v1/finance/payments", headers=headers)
    assert resp.status_code == 403 and resp.json()["code"] == 20002
    resp = client.post("/api/v1/finance/settlements", headers=headers, json={"task_id": 1, "amount": 1})
    assert resp.status_code == 403 and resp.json()["code"] == 20002
    resp = client.post(
        "/api/v1/finance/settlements/1/send-confirm", headers=headers, json={"receiver_id": 1}
    )
    assert resp.status_code == 403 and resp.json()["code"] == 20002
    resp = client.post(
        "/api/v1/finance/settlements/1/confirm", headers=headers, json={"action": "confirm", "note": ""}
    )
    assert resp.status_code == 403 and resp.json()["code"] == 20002


def test_settlement_confirm_flow(client, admin_headers, manager_headers):
    """客户确认：发送（站内通知+留痕）→ 异议 → 重发 → 确认；开票被确认状态门控。"""
    headers = admin_headers
    tid, _task = _finished_task(client, headers)
    r = client.post("/api/v1/finance/settlements", headers=headers, json={"task_id": tid, "amount": 5000}).json()
    assert r["code"] == 0, r
    sid = r["data"]["id"]
    assert r["data"]["confirmStatus"] == ""

    # 收件人 = manager1；不存在用户 → 10002
    users = client.get("/api/v1/admin/users/options", headers=headers, params={"q": "manager1"}).json()["data"]
    assert users and users[0]["label"] == "部门负责人"
    mid = users[0]["id"]
    r = client.post(
        f"/api/v1/finance/settlements/{sid}/send-confirm", headers=headers, json={"receiver_id": 99999999}
    ).json()
    assert r["code"] == 10002

    # 发送 → pending + 收件人站内通知
    r = client.post(
        f"/api/v1/finance/settlements/{sid}/send-confirm",
        headers=headers,
        json={"receiver_id": mid, "note": "请与客户核对"},
    ).json()
    assert r["code"] == 0 and r["data"]["confirmStatus"] == "pending", r
    notifies = client.get(
        "/api/v1/admin/notifies", headers=manager_headers, params={"status": "unread", "size": 100}
    ).json()["data"]["items"]
    hit = [n for n in notifies if n["bizType"] == "settlement" and n["bizId"] == sid]
    assert hit and sid and hit[0]["link"] == "/finance/settlements"

    # 待确认时申请开票 → 80001
    r = client.post(f"/api/v1/finance/settlements/{sid}/apply-invoice", headers=headers, json={}).json()
    assert r["code"] == 80001

    # 异议 → objected（仍可重发）
    r = client.post(
        f"/api/v1/finance/settlements/{sid}/confirm",
        headers=headers,
        json={"action": "object", "note": "客户对金额有异议"},
    ).json()
    assert r["code"] == 0 and r["data"]["confirmStatus"] == "objected", r
    assert r["data"]["confirmBy"] == "系统管理员" and r["data"]["confirmNote"] == "客户对金额有异议"

    # 重发 → pending；确认 → confirmed + confirmTime
    _confirm_settlement(client, headers, sid)
    s = client.get(f"/api/v1/finance/settlements/{sid}", headers=headers).json()["data"]
    assert s["confirmStatus"] == "confirmed" and s["confirmTime"]
    actions = [r["action"] for r in s["confirmRecords"]]
    assert actions == ["send", "object", "send", "confirm"]
    assert s["confirmRecords"][0]["receiverName"] == "部门负责人"

    # 确认后申请开票放行
    r = client.post(f"/api/v1/finance/settlements/{sid}/apply-invoice", headers=headers, json={}).json()
    assert r["code"] == 0 and r["data"]["status"] == "invoicing", r

    # 已进入开票流程：不允许再发送确认/标记确认
    r = client.post(
        f"/api/v1/finance/settlements/{sid}/send-confirm", headers=headers, json={"receiver_id": mid}
    ).json()
    assert r["code"] == 80001

    # 列表可按 confirm_status 过滤
    lst = client.get(
        "/api/v1/finance/settlements", headers=headers, params={"filters[confirm_status]": "confirmed"}
    ).json()["data"]
    assert any(i["id"] == sid for i in lst["items"])


def test_invoice_list_and_export(client, admin_headers, manager_headers):
    """发票列表：审批通过生成 issued 行（带出结算单号/任务号/客户），q 搜索，导出 xlsx。"""
    headers = admin_headers
    tid, _task = _finished_task(client, headers)
    r = client.post("/api/v1/finance/settlements", headers=headers, json={"task_id": tid, "amount": 5000}).json()
    assert r["code"] == 0, r
    sid = r["data"]["id"]
    sno = r["data"]["settlementNo"]
    _confirm_settlement(client, headers, sid)

    # 审批中：发票行尚未生成（审批通过才落行）
    r = client.post(f"/api/v1/finance/settlements/{sid}/apply-invoice", headers=headers, json={}).json()
    assert r["code"] == 0 and r["data"]["status"] == "invoicing", r
    lst = client.get("/api/v1/finance/invoices", headers=headers, params={"q": sno}).json()["data"]
    assert lst["total"] == 0

    # 审批通过 → issued 发票行，关联带出 settlementNo/taskNo/customerName
    todo = _todo_id(client, manager_headers, "invoice", sid)
    r = client.post(
        f"/api/v1/approval/tasks/{todo}/approve", headers=manager_headers, json={"comment": "同意"}
    ).json()
    assert r["code"] == 0, r
    lst = client.get("/api/v1/finance/invoices", headers=headers, params={"q": sno}).json()["data"]
    assert lst["total"] == 1
    inv = lst["items"][0]
    assert inv["status"] == "issued" and inv["invoiceNo"] == ""
    assert inv["settlementId"] == sid and inv["settlementNo"] == sno
    assert inv["taskNo"] and inv["customerName"] and inv["createTime"]

    # 状态页签过滤
    d = client.get("/api/v1/finance/invoices", headers=headers, params={"status": "issued", "q": sno}).json()["data"]
    assert d["total"] == 1
    d = client.get(
        "/api/v1/finance/invoices", headers=headers, params={"status": "approving", "q": sno}
    ).json()["data"]
    assert d["total"] == 0

    # 登记发票号后可按发票号搜索
    inv_no = f"INV-LIST-{uuid.uuid4().hex[:8]}"
    r = client.put(
        f"/api/v1/finance/invoices/{inv['id']}",
        headers=headers,
        json={"invoice_no": inv_no, "amount": 5000, "invoice_date": "2026-10-10"},
    ).json()
    assert r["code"] == 0, r
    d = client.get("/api/v1/finance/invoices", headers=headers, params={"q": inv_no}).json()["data"]
    assert d["total"] == 1 and d["items"][0]["invoiceNo"] == inv_no

    # 导出（GET，xlsx 流）
    resp = client.get("/api/v1/finance/invoices/export", headers=headers)
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.active.max_row >= 2


def test_finance_invalid_state_and_not_found(client, admin_headers):
    headers = admin_headers
    assert client.get("/api/v1/finance/settlements/99999999", headers=headers).json()["code"] == 10002
    # 草稿任务不允许结算 → 80001
    cid = _approved_contract(client, headers)
    r = client.post(
        "/api/v1/task/tasks", headers=headers, json={"task_name": f"草稿任务_{UUID}", "contract_id": cid}
    ).json()
    assert r["code"] == 0, r
    r = client.post(
        "/api/v1/finance/settlements", headers=headers, json={"task_id": r["data"]["id"], "amount": 100}
    ).json()
    assert r["code"] == 80001


def test_settlement_links_and_finance_summary(client, admin_headers, manager_headers):
    """结算单持久化 contract_id/customer_id；三表按合同/客户过滤；合同财务汇总；任务按客户过滤。"""
    headers = admin_headers
    tid, task = _finished_task(client, headers)
    cid, cust_id = task["contractId"], task["customerId"]

    # 合同下拉带 customerId/customerName（建任务自动填充用）
    con = client.get(f"/api/v1/contract/contracts/{cid}", headers=headers).json()["data"]
    opts = client.get(
        "/api/v1/contract/contracts/options", headers=headers, params={"q": con["name"]}
    ).json()["data"]
    opt = next(o for o in opts if o["id"] == cid)
    assert opt["customerId"] == cust_id and opt["customerName"] == task["customerName"]

    # 创建结算单：从任务带出合同/客户关联
    r = client.post("/api/v1/finance/settlements", headers=headers, json={"task_id": tid, "amount": 10000}).json()
    assert r["code"] == 0, r
    sid = r["data"]["id"]
    assert r["data"]["contractId"] == cid and r["data"]["customerId"] == cust_id

    # 结算单列表：filters[contract_id] / filters[customer_id]
    for key, val in (("contract_id", cid), ("customer_id", cust_id)):
        lst = client.get(
            "/api/v1/finance/settlements", headers=headers, params={f"filters[{key}]": val}
        ).json()["data"]
        assert any(i["id"] == sid for i in lst["items"])
        snake = "contractId" if key == "contract_id" else "customerId"
        assert all(i[snake] == val for i in lst["items"])

    # 任务列表：filters[customer_id]
    lst = client.get("/api/v1/task/tasks", headers=headers, params={"filters[customer_id]": cust_id}).json()["data"]
    assert any(i["id"] == tid for i in lst["items"])
    assert all(i["customerId"] == cust_id for i in lst["items"])

    # 开票通过 + 到账 → 发票/到账行
    _confirm_settlement(client, headers, sid)
    assert client.post(f"/api/v1/finance/settlements/{sid}/apply-invoice", headers=headers, json={}).json()["code"] == 0
    todo = _todo_id(client, manager_headers, "invoice", sid)
    assert client.post(
        f"/api/v1/approval/tasks/{todo}/approve", headers=manager_headers, json={"comment": "同意"}
    ).json()["code"] == 0
    r = client.post(
        "/api/v1/finance/payments",
        headers=headers,
        json={"settlement_id": sid, "pay_type": "bank_transfer", "payer_name": "测试付款方",
              "amount": 10000, "paid_at": "2026-10-10"},
    ).json()
    assert r["code"] == 0, r
    assert r["data"]["contractId"] == cid and r["data"]["customerId"] == cust_id

    # 发票/到账列表：两种参数写法均可按合同/客户过滤，行内带出 contractId/customerId
    for params in ({"filters[contract_id]": cid}, {"contract_id": cid}):
        lst = client.get("/api/v1/finance/invoices", headers=headers, params=params).json()["data"]
        assert lst["total"] == 1
        assert lst["items"][0]["contractId"] == cid and lst["items"][0]["customerId"] == cust_id
    lst = client.get(
        "/api/v1/finance/payments", headers=headers, params={"filters[customer_id]": cust_id}
    ).json()["data"]
    assert lst["total"] == 1 and lst["items"][0]["customerId"] == cust_id

    # 合同财务汇总
    d = client.get(f"/api/v1/contract/contracts/{cid}/finance-summary", headers=headers).json()["data"]
    assert d["contractAmount"] == 10000.0
    assert d["settledTotal"] == 10000.0 and d["settlementCount"] == 1
    assert d["invoicedTotal"] == 10000.0 and d["invoiceCount"] == 1
    assert d["paidTotal"] == 10000.0 and d["paymentCount"] == 1

    # 不存在的合同 → 10002
    assert client.get("/api/v1/contract/contracts/99999999/finance-summary", headers=headers).json()["code"] == 10002
