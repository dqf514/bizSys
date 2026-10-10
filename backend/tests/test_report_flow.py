"""报告主链路（HTTP）：上传 → 签批（自动转 PDF 降级）→ 签章 → 任务可结束；门禁 60003 / 重复 61001 / 状态 61002。"""

import io
import uuid

from openpyxl import load_workbook

UUID = uuid.uuid4().hex[:8]

DOCX = b"PK\x03\x04 fake docx content"
PDF = b"%PDF-1.4 fake pdf content"


def _approved_contract(client, headers) -> int:
    suf = uuid.uuid4().hex[:8]
    cust = client.post(
        "/api/v1/crm/customers",
        headers=headers,
        json={"name": f"报告客户_{suf}", "tax_no": f"913401RPT{suf}", "status": "enabled"},
    ).json()
    assert cust["code"] == 0, cust
    con = client.post(
        "/api/v1/contract/contracts",
        headers=headers,
        json={"name": f"报告合同_{suf}", "customer_id": cust["data"]["id"], "amount": 10000},
    ).json()
    cid = con["data"]["id"]
    assert client.post(f"/api/v1/contract/contracts/{cid}/submit", headers=headers, json={}).json()["code"] == 0

    from sqlalchemy import select

    from app.core.db import SessionLocal
    from app.modules.approval import service as approval
    from app.modules.approval.models import ApprovalInstance, ApprovalTask
    from app.modules.auth.models import User

    db = SessionLocal()
    try:
        for _ in range(2):
            target = None
            for t in db.scalars(
                select(ApprovalTask).where(ApprovalTask.status == "pending", ApprovalTask.deleted == 0)
            ):
                inst = db.get(ApprovalInstance, t.instance_id)
                if inst and inst.biz_type == "contract" and inst.biz_id == cid:
                    target = t
                    break
            assert target, "contract instance task not found"
            approval.approve(db, target.id, db.get(User, target.assignee_id), "ok")
            db.commit()
    finally:
        db.close()
    return cid


def _make_task(client, headers, suf: str, need_report: int = 1):
    cid = _approved_contract(client, headers)
    r = client.post(
        "/api/v1/task/tasks",
        headers=headers,
        json={
            "task_name": f"报告任务_{suf}",
            "contract_id": cid,
            "items": [{"name": f"条目A_{suf}", "need_report": need_report}],
        },
    ).json()
    assert r["code"] == 0, r
    return r["data"]["id"], r["data"]["items"][0]


def _upload(client, headers, item_id: int, name="报告.docx", data=DOCX,
            ctype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"):
    return client.post(
        "/api/v1/report/reports",
        headers=headers,
        data={"item_id": str(item_id)},
        files={"file": (name, data, ctype)},
    ).json()


def _act_on_report(biz_action: str, biz_id: int, comment: str = "ok"):
    """直接走引擎审批/驳回报告当前待办（与财务测试同模式）。"""
    from sqlalchemy import select

    from app.core.db import SessionLocal
    from app.modules.approval import service as approval
    from app.modules.approval.models import ApprovalInstance, ApprovalTask
    from app.modules.auth.models import User

    db = SessionLocal()
    try:
        target = None
        for t in db.scalars(select(ApprovalTask).where(ApprovalTask.status == "pending", ApprovalTask.deleted == 0)):
            inst = db.get(ApprovalInstance, t.instance_id)
            if inst and inst.biz_type == "report" and inst.biz_id == biz_id:
                target = t
                break
        assert target, f"report {biz_id} pending task not found"
        user = db.get(User, target.assignee_id)
        if biz_action == "approve":
            approval.approve(db, target.id, user, comment)
        else:
            approval.reject(db, target.id, user, comment)
        db.commit()
    finally:
        db.close()


def _report_status(client, headers, task_id: int, item_id: int) -> str:
    d = client.get(f"/api/v1/task/tasks/{task_id}", headers=headers).json()["data"]
    return next(i["reportStatus"] for i in d["items"] if i["id"] == item_id)


def test_report_full_flow(client, admin_headers):
    headers = admin_headers
    suf = uuid.uuid4().hex[:8]
    tid, item = _make_task(client, headers, suf)
    item_id = item["id"]

    # item-options：格式 {value,label,sub}，label=报告编号
    opts = client.get("/api/v1/report/reports/item-options", headers=headers).json()["data"]
    opt = [o for o in opts if o["value"] == item_id]
    assert opt and opt[0]["label"] == item["reportNo"] and item["name"] in opt[0]["sub"]

    # 非法类型（xlsx 不在报告白名单）→ 10001
    r = _upload(client, headers, item_id, name="报告.xlsx",
                ctype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert r["code"] == 10001

    # 上传 docx → uploaded，条目 reportStatus 同步
    r = _upload(client, headers, item_id)
    assert r["code"] == 0, r
    rid = r["data"]["id"]
    assert r["data"]["reportNo"] == item["reportNo"] and r["data"]["status"] == "uploaded"
    assert _report_status(client, headers, tid, item_id) == "uploaded"

    # 重复上传 → 61001；该条目从下拉消失
    assert _upload(client, headers, item_id)["code"] == 61001
    opts = client.get("/api/v1/report/reports/item-options", headers=headers).json()["data"]
    assert all(o["value"] != item_id for o in opts)

    # 任务跑到 finished，close 被门禁拦住（60003，消息含报告编号）
    for action in ("submit", "start", "start", "finish"):
        assert client.post(f"/api/v1/task/tasks/{tid}/{action}", headers=headers).json()["code"] == 0
    r = client.post(f"/api/v1/task/tasks/{tid}/close", headers=headers).json()
    assert r["code"] == 60003 and item["reportNo"] in r["message"], r

    # 提交审批 → approving；重复提交 → 61002
    r = client.post(f"/api/v1/report/reports/{rid}/submit", headers=headers, json={}).json()
    assert r["code"] == 0 and r["data"]["status"] == "approving"
    assert client.post(f"/api/v1/report/reports/{rid}/submit", headers=headers, json={}).json()["code"] == 61002
    assert _report_status(client, headers, tid, item_id) == "approving"

    # 批准 → approved + 自动转 PDF 降级（本机无 soffice：pdf_file_id=原件 + remark 标注）
    _act_on_report("approve", rid)
    d = client.get(f"/api/v1/report/reports/{rid}", headers=headers).json()["data"]
    assert d["status"] == "approved" and d["pdfReady"]
    assert "未安装转换器" in d["remark"], d["remark"]
    assert _report_status(client, headers, tid, item_id) == "approved"
    assert any(rec["action"] == "agree" for rec in d["timeline"])

    # 未 approved 不能签章已由上面状态保证；签章 → sealed
    r = client.post(f"/api/v1/report/reports/{rid}/seal", headers=headers, json={}).json()
    assert r["code"] == 0 and r["data"]["status"] == "sealed" and r["data"]["sealTime"]
    assert client.post(f"/api/v1/report/reports/{rid}/seal", headers=headers, json={}).json()["code"] == 61002
    assert _report_status(client, headers, tid, item_id) == "sealed"

    # 签章完成后任务可结束
    r = client.post(f"/api/v1/task/tasks/{tid}/close", headers=headers).json()
    assert r["code"] == 0 and r["data"]["status"] == "closed", r

    # 下载原件 / 下载PDF（降级=原件）/ 导出
    resp = client.get(f"/api/v1/report/reports/{rid}/download", headers=headers)
    assert resp.status_code == 200 and resp.content == DOCX
    resp = client.get(f"/api/v1/report/reports/{rid}/download", headers=headers, params={"pdf": 1})
    assert resp.status_code == 200 and resp.content == DOCX
    resp = client.get("/api/v1/report/reports/export", headers=headers)
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.active.max_row >= 2

    # 列表检索：状态页签 + 关键词（任务号）
    lst = client.get("/api/v1/report/reports", headers=headers, params={"status": "sealed"}).json()["data"]
    assert any(i["id"] == rid for i in lst["items"])
    lst = client.get("/api/v1/report/reports", headers=headers, params={"q": f"报告任务_{suf}"}).json()["data"]
    assert lst["total"] == 1 and lst["items"][0]["sealBy"]


def test_report_reject_resubmit_flow(client, admin_headers):
    """驳回 → 可重新提交；撤回 → 回退 uploaded。report_status 同步。"""
    headers = admin_headers
    suf = uuid.uuid4().hex[:8]
    tid, item = _make_task(client, headers, suf)
    item_id = item["id"]
    r = _upload(client, headers, item_id, name="报告.pdf", data=PDF, ctype="application/pdf")
    assert r["code"] == 0, r
    rid = r["data"]["id"]

    assert client.post(f"/api/v1/report/reports/{rid}/submit", headers=headers, json={}).json()["code"] == 0

    # 发起人撤回（尚无人处理）→ uploaded
    from sqlalchemy import select

    from app.core.db import SessionLocal
    from app.modules.approval import service as approval
    from app.modules.auth.models import User

    db = SessionLocal()
    try:
        inst_id = client.get(f"/api/v1/report/reports/{rid}", headers=headers).json()["data"]["approvalInstanceId"]
        admin = db.scalar(select(User).where(User.username == "admin"))
        approval.withdraw(db, inst_id, admin)
        db.commit()
    finally:
        db.close()
    d = client.get(f"/api/v1/report/reports/{rid}", headers=headers).json()["data"]
    assert d["status"] == "uploaded"
    assert _report_status(client, headers, tid, item_id) == "uploaded"

    # 再提交 → 驳回 → rejected；驳回后仍可再提交
    assert client.post(f"/api/v1/report/reports/{rid}/submit", headers=headers, json={}).json()["code"] == 0
    _act_on_report("reject", rid, "内容需修改")
    d = client.get(f"/api/v1/report/reports/{rid}", headers=headers).json()["data"]
    assert d["status"] == "rejected"
    assert _report_status(client, headers, tid, item_id) == "rejected"

    assert client.post(f"/api/v1/report/reports/{rid}/submit", headers=headers, json={}).json()["code"] == 0
    _act_on_report("approve", rid)
    d = client.get(f"/api/v1/report/reports/{rid}", headers=headers).json()["data"]
    # 原件即 PDF：直接复用，无降级标注
    assert d["status"] == "approved" and d["pdfReady"] and "未安装转换器" not in d["remark"]
    assert client.post(f"/api/v1/report/reports/{rid}/seal", headers=headers, json={}).json()["code"] == 0
    resp = client.get(f"/api/v1/report/reports/{rid}/download", headers=headers, params={"pdf": 1})
    assert resp.status_code == 200 and resp.content == PDF


def test_close_guard_without_report(client, admin_headers):
    """need_report=1 但从未上传报告：close 直接 60003。"""
    headers = admin_headers
    suf = uuid.uuid4().hex[:8]
    tid, item = _make_task(client, headers, suf)
    for action in ("submit", "start", "start", "finish"):
        assert client.post(f"/api/v1/task/tasks/{tid}/{action}", headers=headers).json()["code"] == 0
    r = client.post(f"/api/v1/task/tasks/{tid}/close", headers=headers).json()
    assert r["code"] == 60003 and item["name"] in r["message"], r


def test_report_list_contract_filter(client, admin_headers):
    """报告列表 contract_id 过滤（经任务关联），行内带出 contractId/contractNo；item-options 支持 task_id。"""
    headers = admin_headers
    suf = uuid.uuid4().hex[:8]
    cid = _approved_contract(client, headers)
    r = client.post(
        "/api/v1/task/tasks",
        headers=headers,
        json={
            "task_name": f"合同过滤任务_{suf}",
            "contract_id": cid,
            "items": [{"name": f"条目B_{suf}", "need_report": 1}],
        },
    ).json()
    assert r["code"] == 0, r
    tid = r["data"]["id"]
    item = r["data"]["items"][0]

    # item-options：task_id 限定到本任务 / 其他任务为空
    opts = client.get(
        "/api/v1/report/reports/item-options", headers=headers, params={"task_id": tid}
    ).json()["data"]
    assert any(o["value"] == item["id"] for o in opts)
    assert client.get(
        "/api/v1/report/reports/item-options", headers=headers, params={"task_id": tid + 10**6}
    ).json()["data"] == []

    r = _upload(client, headers, item["id"])
    assert r["code"] == 0, r
    rid = r["data"]["id"]

    # contract_id 直参 / filters[contract_id] 两种写法等价
    for params in ({"contract_id": cid}, {"filters[contract_id]": cid}):
        lst = client.get("/api/v1/report/reports", headers=headers, params=params).json()["data"]
        row = next(i for i in lst["items"] if i["id"] == rid)
        assert row["contractId"] == cid and row["contractNo"]
        assert all(i["contractId"] == cid for i in lst["items"])

    # customer_id 同样支持两种写法（经任务关联过滤）
    cust_id = client.get(f"/api/v1/contract/contracts/{cid}", headers=headers).json()["data"]["customerId"]
    for params in ({"customer_id": cust_id}, {"filters[customer_id]": cust_id}):
        lst = client.get("/api/v1/report/reports", headers=headers, params=params).json()["data"]
        assert any(i["id"] == rid for i in lst["items"])
    lst = client.get(
        "/api/v1/report/reports", headers=headers, params={"filters[customer_id]": cust_id + 10**6}
    ).json()["data"]
    assert all(i["id"] != rid for i in lst["items"])

    # 无关合同查不到该报告
    lst = client.get(
        "/api/v1/report/reports", headers=headers, params={"contract_id": cid + 10**6}
    ).json()["data"]
    assert all(i["id"] != rid for i in lst["items"])
