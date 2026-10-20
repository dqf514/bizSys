"""合同主链路（HTTP）：客户 → 合同 → 提交 → 两级审批 → 已通过 → 导出。"""

import io
import uuid

from openpyxl import load_workbook

UUID = uuid.uuid4().hex[:8]


def _create_customer(client, headers):
    suf = uuid.uuid4().hex[:8]  # 每次调用独立后缀：测试库数据跨用例/跨运行持久
    r = client.post(
        "/api/v1/crm/customers",
        headers=headers,
        json={
            "name": f"测试客户_{suf}",
            "tax_no": f"913401TEST{suf}",
            "customer_type": "enterprise",
            "grade": "A",
            "status": "enabled",
        },
    ).json()
    assert r["code"] == 0, r
    return r["data"]


def test_contract_full_flow(client, admin_headers, manager_headers, leader_headers):
    cust = _create_customer(client, admin_headers)

    # 税号重复 40001
    r = client.post(
        "/api/v1/crm/customers", headers=admin_headers, json={"name": "dup", "tax_no": cust["taxNo"]}
    ).json()
    assert r["code"] == 40001

    # 建合同
    r = client.post(
        "/api/v1/contract/contracts",
        headers=admin_headers,
        json={
            "name": f"合同_{UUID}",
            "customer_id": cust["id"],
            "contract_type": "service",
            "amount": 50000,
            "start_date": "2026-10-01",
            "end_date": "2027-09-30",
        },
    ).json()
    assert r["code"] == 0, r
    cid = r["data"]["id"]
    # F3：draft 不占正式 HT 编号，提交审批时发放
    assert r["data"]["contractNo"] == ""

    # 提交审批
    r = client.post(f"/api/v1/contract/contracts/{cid}/submit", headers=admin_headers, json={"comment": "审批"}).json()
    assert r["code"] == 0 and r["data"]["status"] == "approving"
    assert r["data"]["contractNo"].startswith("HT")

    # 审批中禁编辑 50002
    r = client.put(
        f"/api/v1/contract/contracts/{cid}",
        headers=admin_headers,
        json={"name": "x", "customer_id": cust["id"], "version": 1},
    ).json()
    assert r["code"] == 50002

    # manager1 同意
    page = client.get("/api/v1/approval/tasks", headers=manager_headers, params={"tab": "todo"}).json()
    todo = [i for i in page["data"]["items"] if i["bizId"] == cid]
    assert todo
    r = client.post(
        f"/api/v1/approval/tasks/{todo[0]['id']}/approve", headers=manager_headers, json={"comment": "ok"}
    ).json()
    assert r["code"] == 0 and r["data"]["status"] == "running"

    # leader1 同意
    page = client.get("/api/v1/approval/tasks", headers=leader_headers, params={"tab": "todo"}).json()
    todo = [i for i in page["data"]["items"] if i["bizId"] == cid]
    assert todo
    r = client.post(
        f"/api/v1/approval/tasks/{todo[0]['id']}/approve", headers=leader_headers, json={"comment": "ok"}
    ).json()
    assert r["code"] == 0 and r["data"]["status"] == "approved"

    # 合同已通过 + 时间线
    r = client.get(f"/api/v1/contract/contracts/{cid}", headers=admin_headers).json()
    assert r["data"]["status"] == "approved"
    assert [x["action"] for x in r["data"]["timeline"]] == ["start", "agree", "agree"]

    # options 可见
    opts = client.get("/api/v1/contract/contracts/options", headers=admin_headers).json()["data"]
    assert any(o["id"] == cid for o in opts)

    # 导出
    resp = client.post("/api/v1/contract/contracts/export", headers=admin_headers)
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.active.max_row >= 2


def test_withdraw_then_edit_and_resubmit(client, admin_headers):
    """提交 → 撤回 → 合同回 draft，可再编辑/再提交（修复前卡在 approving）。"""
    cust = _create_customer(client, admin_headers)
    r = client.post(
        "/api/v1/contract/contracts",
        headers=admin_headers,
        json={"name": f"撤回合同_{UUID}", "customer_id": cust["id"], "amount": 8000},
    ).json()
    assert r["code"] == 0, r
    cid = r["data"]["id"]
    r = client.post(f"/api/v1/contract/contracts/{cid}/submit", headers=admin_headers, json={}).json()
    assert r["code"] == 0 and r["data"]["status"] == "approving"
    inst_id = r["data"]["flowInstanceId"]

    r = client.post(f"/api/v1/approval/instances/{inst_id}/withdraw", headers=admin_headers, json={}).json()
    assert r["code"] == 0 and r["data"]["status"] == "withdrawn"
    r = client.get(f"/api/v1/contract/contracts/{cid}", headers=admin_headers).json()
    assert r["data"]["status"] == "draft"

    # 可再编辑
    r = client.put(
        f"/api/v1/contract/contracts/{cid}",
        headers=admin_headers,
        json={"name": f"撤回合同改_{UUID}", "customer_id": cust["id"], "amount": 9000, "version": r["data"]["version"]},
    ).json()
    assert r["code"] == 0, r

    # 可再提交
    r = client.post(f"/api/v1/contract/contracts/{cid}/submit", headers=admin_headers, json={}).json()
    assert r["code"] == 0 and r["data"]["status"] == "approving"


def test_attachment_upload_validation(client, admin_headers):
    """附件上传：扩展名白名单 / 空文件 → 业务错误码而非 500。"""
    cust = _create_customer(client, admin_headers)
    r = client.post(
        "/api/v1/contract/contracts",
        headers=admin_headers,
        json={"name": f"附件合同_{UUID}", "customer_id": cust["id"], "amount": 100},
    ).json()
    cid = r["data"]["id"]

    r = client.post(
        f"/api/v1/contract/contracts/{cid}/attachments",
        headers=admin_headers,
        files={"file": ("evil.exe", io.BytesIO(b"MZ"), "application/octet-stream")},
    ).json()
    assert r["code"] == 10001

    r = client.post(
        f"/api/v1/contract/contracts/{cid}/attachments",
        headers=admin_headers,
        files={"file": ("a.pdf", io.BytesIO(b"%PDF-1.4 fake"), "application/pdf")},
    ).json()
    assert r["code"] == 0, r
    assert r["data"]["fileName"] == "a.pdf"


def test_attachment_delete_state_guard(client, admin_headers):
    """附件删除：draft 可删（逻辑删，详情不再返回）；approving 拒 50002。"""
    cust = _create_customer(client, admin_headers)
    r = client.post(
        "/api/v1/contract/contracts",
        headers=admin_headers,
        json={"name": f"附件删除_{UUID}", "customer_id": cust["id"], "amount": 100},
    ).json()
    cid = r["data"]["id"]

    # draft：上传后可删
    r = client.post(
        f"/api/v1/contract/contracts/{cid}/attachments",
        headers=admin_headers,
        files={"file": ("draft.pdf", io.BytesIO(b"%PDF-1.4 fake"), "application/pdf")},
    ).json()
    assert r["code"] == 0, r
    fid = r["data"]["id"]
    r = client.delete(f"/api/v1/contract/attachments/{fid}", headers=admin_headers).json()
    assert r["code"] == 0, r
    detail = client.get(f"/api/v1/contract/contracts/{cid}", headers=admin_headers).json()["data"]
    assert all(a["id"] != fid for a in detail["attachments"])
    # 再删一次 → 404 业务码
    r = client.delete(f"/api/v1/contract/attachments/{fid}", headers=admin_headers).json()
    assert r["code"] == 10002

    # approving：拒绝删除
    r = client.post(
        f"/api/v1/contract/contracts/{cid}/attachments",
        headers=admin_headers,
        files={"file": ("lock.pdf", io.BytesIO(b"%PDF-1.4 fake"), "application/pdf")},
    ).json()
    fid2 = r["data"]["id"]
    assert client.post(f"/api/v1/contract/contracts/{cid}/submit", headers=admin_headers, json={}).json()["code"] == 0
    r = client.delete(f"/api/v1/contract/attachments/{fid2}", headers=admin_headers).json()
    assert r["code"] == 50002


def test_contract_draft_minimal_fields(client, admin_headers):
    """保存草稿只需名称+客户：日期/金额不强制；空串日期宽容解析为 None（不再 422）。"""
    cust = _create_customer(client, admin_headers)
    r = client.post(
        "/api/v1/contract/contracts",
        headers=admin_headers,
        json={"name": f"草稿合同_{UUID}", "customer_id": cust["id"], "start_date": "", "end_date": ""},
    ).json()
    assert r["code"] == 0, r
    assert r["data"]["status"] == "draft"
    assert r["data"]["contractNo"] == ""
    assert r["data"]["startDate"] == "" and r["data"]["endDate"] == ""

    # 提交审批才做完整校验：无金额 → 10001
    cid = r["data"]["id"]
    r = client.post(f"/api/v1/contract/contracts/{cid}/submit", headers=admin_headers, json={}).json()
    assert r["code"] == 10001, r
