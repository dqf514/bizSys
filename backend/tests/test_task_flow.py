"""任务模块（HTTP）：合同校验 60002 / 状态机 60001 / 编号 / 计数 / 导出。"""

import io
import uuid

from openpyxl import load_workbook

UUID = uuid.uuid4().hex[:8]


def _approved_contract(client, headers):
    suf = uuid.uuid4().hex[:8]  # 每次调用独立后缀：测试库数据跨用例持久
    cust = client.post(
        "/api/v1/crm/customers",
        headers=headers,
        json={"name": f"任务客户_{suf}", "tax_no": f"913401TSK{suf}", "status": "enabled"},
    ).json()
    assert cust["code"] == 0, cust
    con = client.post(
        "/api/v1/contract/contracts",
        headers=headers,
        json={"name": f"任务合同_{suf}", "customer_id": cust["data"]["id"], "amount": 10000},
    ).json()
    cid = con["data"]["id"]
    assert client.post(f"/api/v1/contract/contracts/{cid}/submit", headers=headers, json={}).json()["code"] == 0

    from sqlalchemy import select

    from app.core.db import SessionLocal
    from app.modules.approval import service as approval
    from app.modules.approval.models import ApprovalTask

    db = SessionLocal()
    try:
        for _ in range(2):  # 两级依次同意
            rows = db.scalars(
                select(ApprovalTask).where(
                    ApprovalTask.instance_id != 0, ApprovalTask.status == "pending", ApprovalTask.deleted == 0
                )
            ).all()
            mine = [t for t in rows]
            assert mine, "no pending approval task"
            # 找到本合同实例的待办
            from app.modules.approval.models import ApprovalInstance

            target = None
            for t in mine:
                inst = db.get(ApprovalInstance, t.instance_id)
                if inst and inst.biz_type == "contract" and inst.biz_id == cid:
                    target = t
                    break
            assert target, "contract instance task not found"
            from app.modules.auth.models import User

            assignee = db.get(User, target.assignee_id)
            approval.approve(db, target.id, assignee, "ok")
            db.commit()
    finally:
        db.close()
    return cid


def _approve_report(rid: int):
    """批准报告签批当前待办（一级：manager1）。"""
    from sqlalchemy import select

    from app.core.db import SessionLocal
    from app.modules.approval import service as approval
    from app.modules.approval.models import ApprovalInstance, ApprovalTask
    from app.modules.auth.models import User

    db = SessionLocal()
    try:
        target = None
        for t in db.scalars(select(ApprovalTask).where(ApprovalTask.status == "pending", ApprovalTask.deleted == 0)):
            inst = db.get(ApprovalInstance, t.instance_id)
            if inst and inst.biz_type == "report" and inst.biz_id == rid:
                target = t
                break
        assert target, "report pending task not found"
        approval.approve(db, target.id, db.get(User, target.assignee_id), "ok")
        db.commit()
    finally:
        db.close()


def test_task_state_machine_and_errors(client, admin_headers):
    headers = admin_headers
    approved_id = _approved_contract(client, headers)

    # 未通过合同 60002（客户自建，避免依赖种子数据 id）
    suf = uuid.uuid4().hex[:8]
    cust = client.post(
        "/api/v1/crm/customers",
        headers=headers,
        json={"name": f"草稿客户_{suf}", "tax_no": f"913401TSK{suf}", "status": "enabled"},
    ).json()
    assert cust["code"] == 0, cust
    draft = client.post(
        "/api/v1/contract/contracts",
        headers=headers,
        json={"name": f"草稿_{UUID}", "customer_id": cust["data"]["id"], "amount": 1},
    ).json()
    r = client.post(
        "/api/v1/task/tasks", headers=headers, json={"task_name": "x", "contract_id": draft["data"]["id"]}
    ).json()
    assert r["code"] == 60002

    # 建任务（条目带报告编号，样品自动编号）
    r = client.post(
        "/api/v1/task/tasks",
        headers=headers,
        json={
            "task_name": f"任务_{UUID}",
            "task_type": "wind_tunnel",
            "contract_id": approved_id,
            "biz_block": "wind",
            "items": [{"name": "条目1", "need_report": 1}],
            "samples": [{"name": "样车"}],
        },
    ).json()
    assert r["code"] == 0, r
    tid = r["data"]["id"]
    # 草稿不占正式编号：任务号/报告号/样品号均待提交时发放
    assert r["data"]["taskNo"] == ""
    assert r["data"]["items"][0]["reportNo"] == ""
    assert r["data"]["samples"][0]["sampleNo"] == ""
    assert r["data"]["customerName"]

    # 非法流转 60001
    r = client.post(f"/api/v1/task/tasks/{tid}/finish", headers=headers).json()
    assert r["code"] == 60001

    # 全链路流转到 in_progress；条目 need_report=1 未签章 → finish 门禁 60003
    r = client.post(f"/api/v1/task/tasks/{tid}/submit", headers=headers).json()
    assert r["code"] == 0 and r["data"]["status"] == "in_progress", r
    r = client.post(f"/api/v1/task/tasks/{tid}/finish", headers=headers).json()
    assert r["code"] == 60003, r

    # 提交时已发放正式编号
    detail = client.get(f"/api/v1/task/tasks/{tid}", headers=headers).json()["data"]
    tno = detail["taskNo"]
    assert tno.startswith("RW")
    assert detail["items"][0]["reportNo"].startswith("BG")
    assert detail["samples"][0]["sampleNo"] == f"{tno}-01"
    item_id = detail["items"][0]["id"]
    # 补齐报告链路（上传 → 签批 → 批准 → 签章）后 finish 放行
    r = client.post(
        "/api/v1/report/reports",
        headers=headers,
        data={"item_id": str(item_id)},
        files={"file": ("报告.pdf", b"%PDF-1.4 fake", "application/pdf")},
    ).json()
    assert r["code"] == 0, r
    rid = r["data"]["id"]
    assert client.post(f"/api/v1/report/reports/{rid}/submit", headers=headers, json={}).json()["code"] == 0
    _approve_report(rid)
    assert client.post(f"/api/v1/report/reports/{rid}/seal", headers=headers, json={}).json()["code"] == 0
    r = client.post(f"/api/v1/task/tasks/{tid}/finish", headers=headers).json()
    assert r["code"] == 0 and r["data"]["status"] == "finished", r

    # 终态禁改
    r = client.put(f"/api/v1/task/tasks/{tid}", headers=headers, json={"task_name": "y", "version": 1}).json()
    assert r["code"] == 60001

    # 计数 + 检索 + 导出
    counts = client.get("/api/v1/task/tasks/counts", headers=headers).json()["data"]
    assert counts["finished"] >= 1 and counts["all"] >= 1
    lst = client.get("/api/v1/task/tasks", headers=headers, params={"q": f"任务_{UUID}"}).json()["data"]
    assert lst["total"] == 1
    assert lst["items"][0]["taskTypeLabel"]
    resp = client.post("/api/v1/task/tasks/export", headers=headers)
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.active.max_row >= 2


def test_sample_no_unique_after_delete(client, admin_headers):
    """逻辑删除样品后再新增：sample_no 取历史最大位次+1，不撞唯一索引（修复前 500）。"""
    headers = admin_headers
    approved_id = _approved_contract(client, headers)
    r = client.post(
        "/api/v1/task/tasks",
        headers=headers,
        json={"task_name": f"样品_{UUID}", "contract_id": approved_id, "samples": [{"name": "样车A"}]},
    ).json()
    assert r["code"] == 0, r
    tid = r["data"]["id"]
    sid = r["data"]["samples"][0]["id"]
    # 草稿无编号；提交后才发放
    assert r["data"]["samples"][0]["sampleNo"] == ""
    assert client.post(f"/api/v1/task/tasks/{tid}/submit", headers=headers).json()["code"] == 0
    detail = client.get(f"/api/v1/task/tasks/{tid}", headers=headers).json()["data"]
    tno = detail["taskNo"]
    assert tno.startswith("RW")
    assert detail["samples"][0]["sampleNo"] == f"{tno}-01"

    assert client.delete(f"/api/v1/task/samples/{sid}", headers=headers).json()["code"] == 0
    r = client.post(f"/api/v1/task/tasks/{tid}/samples", headers=headers, json={"name": "样车B"}).json()
    assert r["code"] == 0, r
    assert r["data"]["sampleNo"] == f"{tno}-02"


def test_sample_no_assigned_on_submit_batch(client, admin_headers):
    """草稿同步（sync_samples）不编号；提交时按当前样品顺序统一发放且不撞号。"""
    headers = admin_headers
    approved_id = _approved_contract(client, headers)
    r = client.post(
        "/api/v1/task/tasks",
        headers=headers,
        json={
            "task_name": f"批量_{UUID}",
            "contract_id": approved_id,
            "samples": [{"name": "样车A"}, {"name": "样车B"}],
        },
    ).json()
    assert r["code"] == 0, r
    tid = r["data"]["id"]
    sid_a = r["data"]["samples"][0]["id"]
    assert all(s["sampleNo"] == "" for s in r["data"]["samples"])

    # 草稿期删除 A 并同步新增 C：仍全部无编号
    assert client.delete(f"/api/v1/task/samples/{sid_a}", headers=headers).json()["code"] == 0
    keep = r["data"]["samples"][1]
    r = client.put(
        f"/api/v1/task/tasks/{tid}",
        headers=headers,
        json={"task_name": f"批量_{UUID}", "samples": [keep, {"name": "样车C"}]},
    ).json()
    assert r["code"] == 0, r
    assert all(s["sampleNo"] == "" for s in r["data"]["samples"])

    # 提交 → 统一编号（A 从未占号，B/C 依次 -01/-02）
    assert client.post(f"/api/v1/task/tasks/{tid}/submit", headers=headers).json()["code"] == 0
    detail = client.get(f"/api/v1/task/tasks/{tid}", headers=headers).json()["data"]
    tno = detail["taskNo"]
    nos = [s["sampleNo"] for s in detail["samples"]]
    assert nos == [f"{tno}-01", f"{tno}-02"], nos


def test_task_detail_logs(client, admin_headers):
    """任务详情 logs：样品登记/出入库流水按时间倒序返回，含样品编号/动作/操作人。"""
    headers = admin_headers
    approved_id = _approved_contract(client, headers)
    r = client.post(
        "/api/v1/task/tasks",
        headers=headers,
        json={"task_name": f"日志_{UUID}", "contract_id": approved_id},
    ).json()
    assert r["code"] == 0, r
    tid = r["data"]["id"]

    # 提交发号后再补登样品（写「初始登记」流水），样品即带编号
    assert client.post(f"/api/v1/task/tasks/{tid}/submit", headers=headers).json()["code"] == 0
    r = client.post(f"/api/v1/task/tasks/{tid}/samples", headers=headers, json={"name": "样车X"}).json()
    assert r["code"] == 0, r
    sid = r["data"]["id"]
    sno = r["data"]["sampleNo"]

    assert client.post(f"/api/v1/task/samples/{sid}/out", headers=headers).json()["code"] == 0
    assert client.post(f"/api/v1/task/samples/{sid}/in", headers=headers).json()["code"] == 0

    r = client.get(f"/api/v1/task/tasks/{tid}", headers=headers).json()
    assert r["code"] == 0, r
    logs = r["data"]["logs"]
    assert len(logs) >= 3, logs
    # 倒序：最新的入库在最前
    assert [x["action"] for x in logs[:3]] == ["in", "out", "in"]
    assert all(x["sampleNo"] == sno for x in logs)
    assert all(x["operatorName"] for x in logs)
    assert logs[-1]["remark"] == "初始登记"


def test_invalid_date_filter_is_biz_error(client, admin_headers):
    """非法日期区间 → 10001 业务错误，不再 500。"""
    r = client.get(
        "/api/v1/task/tasks",
        headers=admin_headers,
        params={"range[plan_finish_date]": "not-a-date|2026-12-31"},
    )
    data = r.json()
    assert data["code"] == 10001
    assert r.status_code == 200


def test_task_draft_empty_date_tolerant(client, admin_headers):
    """日期控件清空回传空串：草稿保存宽容解析为 None（planFinishDate 空串 422 修复）。"""
    r = client.post(
        "/api/v1/task/tasks",
        headers=admin_headers,
        json={
            "task_name": f"空日期_{UUID}",
            "plan_finish_date": "",
            "samples": [{"name": "样车", "sample_date": ""}],
        },
    ).json()
    assert r["code"] == 0, r
    assert r["data"]["planFinishDate"] == ""
    assert r["data"]["samples"][0]["sampleDate"] == ""
    # 合同空串日期同样宽容（走同一 OptionalDate）
    suf = uuid.uuid4().hex[:8]
    cust = client.post(
        "/api/v1/crm/customers",
        headers=admin_headers,
        json={"name": f"空日期客户_{suf}", "tax_no": f"913401TSK{suf}", "status": "enabled"},
    ).json()
    assert cust["code"] == 0, cust
    r = client.post(
        "/api/v1/contract/contracts",
        headers=admin_headers,
        json={"name": f"空日期合同_{UUID}", "customer_id": cust["data"]["id"], "start_date": "", "end_date": ""},
    ).json()
    assert r["code"] == 0, r
